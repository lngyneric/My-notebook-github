from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_presentation():
    prs = Presentation()

    # Define a simple theme color (Blue)
    theme_color = RGBColor(0, 112, 192)

    def add_title_slide(prs, title_text, subtitle_text):
        slide_layout = prs.slide_layouts[0] # Title Slide
        slide = prs.slides.add_slide(slide_layout)
        title = slide.shapes.title
        subtitle = slide.placeholders[1]
        
        title.text = title_text
        subtitle.text = subtitle_text
        
        # Formatting
        title.text_frame.paragraphs[0].font.color.rgb = theme_color
        title.text_frame.paragraphs[0].font.bold = True
        # Ensure Chinese font support if possible, usually defaults work but explicit is better if needed
        # For simplicity, relying on default font fallback

    def add_bullet_slide(prs, title_text, main_points):
        slide_layout = prs.slide_layouts[1] # Title and Content
        slide = prs.slides.add_slide(slide_layout)
        
        title = slide.shapes.title
        title.text = title_text
        title.text_frame.paragraphs[0].font.color.rgb = theme_color
        
        content = slide.placeholders[1]
        tf = content.text_frame
        tf.clear() # Clear default
        
        for point in main_points:
            p = tf.add_paragraph()
            p.text = point['text']
            p.level = point.get('level', 0)
            if point.get('bold', False):
                p.font.bold = True

    # 1. Title Slide
    add_title_slide(prs, 'Trae IDE Agent Skills 案例研究', '自动化入职与管理工作流')

    # 2. Project Overview
    add_bullet_slide(prs, '项目概览', [
        {'text': '目标', 'bold': True},
        {'text': '利用 AI Agent 实现新员工入职、组织管理和团队赋能的自动化。', 'level': 1},
        {'text': '关键技术', 'bold': True},
        {'text': 'Trae IDE, Claude Agent Skills, Python (Pandas), HTML/JS/Tailwind。', 'level': 1}
    ])

    # 3. Skill 1: Intelligent Excel Processing
    add_bullet_slide(prs, 'Skill 1: 智能 Excel 处理', [
        {'text': '使用的 Skill: xlsx / RunCommand', 'bold': True},
        {'text': '挑战：管理复杂的组织架构图和培训数据。', 'level': 0},
        {'text': '采取的行动：', 'level': 0},
        {'text': '分析 "中欧总表.xlsx" 提取培训进度 (邹锋静: 4%)。', 'level': 1},
        {'text': '使用正则批量将敏感人员姓名替换为 "待定"。', 'level': 1},
        {'text': '将表格数据转换为 Mermaid 结构图。', 'level': 1},
        {'text': '成果：自动化的数据清洗与可视化准备。', 'level': 1}
    ])

    # 4. Skill 2: Workflow Automation
    add_bullet_slide(prs, 'Skill 2: 工作流自动化', [
        {'text': '使用的 Skill: outlook-sender', 'bold': True},
        {'text': '挑战：报告需要手动发送邮件通知。', 'level': 0},
        {'text': '采取的行动：', 'level': 0},
        {'text': '创建自定义 Skill 通过 PowerShell 调用本地 Outlook。', 'level': 1},
        {'text': '设计脚本发送带附件的邮件及保存草稿。', 'level': 1},
        {'text': '成果：标准化的通信能力 (概念验证)。', 'level': 1}
    ])

    # 5. Skill 3: Frontend Data Visualization
    add_bullet_slide(prs, 'Skill 3: 前端数据可视化', [
        {'text': '使用的 Skill: frontend-design', 'bold': True},
        {'text': '挑战：纯文本的状态报告难以阅读。', 'level': 0},
        {'text': '采取的行动：', 'level': 0},
        {'text': '将 Markdown 状态报告转化为交互式 HTML 仪表盘。', 'level': 1},
        {'text': '使用 Chart.js 实现实时进度图表。', 'level': 1},
        {'text': '添加交互功能：职级检测与权重调整。', 'level': 1},
        {'text': '成果：专业、响应式的入职状态看板。', 'level': 1}
    ])

    # 6. Skill 4: Knowledge Enablement
    add_bullet_slide(prs, 'Skill 4: 知识赋能', [
        {'text': '使用的 Skill: enjoy-bot-community-guide', 'bold': True},
        {'text': '挑战：团队缺乏使用 AI 社区资源的最佳实践。', 'level': 0},
        {'text': '采取的行动：', 'level': 0},
        {'text': '将社区指南封装为可检索的 Skill。', 'level': 1},
        {'text': '定义 "A-B-C" 使用策略 (场景、结构、行动)。', 'level': 1},
        {'text': '标准化 "高质量提问模板"。', 'level': 1},
        {'text': '成果：可扩展的知识传递机制。', 'level': 1}
    ])

    # 7. Summary
    add_bullet_slide(prs, '总结', [
        {'text': 'Agent Skills 的价值：', 'bold': True},
        {'text': '结构化：将临时任务转化为可复用的模块。', 'level': 1},
        {'text': '自动化：减少数据处理和报告中的手动工作。', 'level': 1},
        {'text': '交互式：从静态文本转向动态可视化工具。', 'level': 1},
        {'text': '本项目展示了 Trae IDE Agents 如何作为全栈工程师处理端到端的工作流。', 'level': 0}
    ])

    prs.save('Trae_Agent_Skills_Case_Study_CN.pptx')
    print('Presentation saved successfully.')

create_presentation()
