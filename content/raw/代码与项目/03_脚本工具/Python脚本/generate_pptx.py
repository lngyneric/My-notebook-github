from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_presentation():
    prs = Presentation()

    # Define a simple theme color (Blue)
    theme_color = RGBColor(0, 112, 192)

    def add_title_slide(prs, title_text, subtitle_text):
        slide_layout = prs.slide_layouts[0] # Title Slide
        slide = prs.slides.add_slide(slide_layout)
        title = slide.shapes.title
        subtitle = slide.placeholders[1]
        
        title.text = title_text
        subtitle.text = subtitle_text
        
        # Formatting
        title.text_frame.paragraphs[0].font.color.rgb = theme_color
        title.text_frame.paragraphs[0].font.bold = True

    def add_bullet_slide(prs, title_text, main_points):
        slide_layout = prs.slide_layouts[1] # Title and Content
        slide = prs.slides.add_slide(slide_layout)
        
        title = slide.shapes.title
        title.text = title_text
        title.text_frame.paragraphs[0].font.color.rgb = theme_color
        
        content = slide.placeholders[1]
        tf = content.text_frame
        tf.clear() # Clear default
        
        for point in main_points:
            p = tf.add_paragraph()
            p.text = point['text']
            p.level = point.get('level', 0)
            if point.get('bold', False):
                p.font.bold = True

    # 1. Title Slide
    add_title_slide(prs, 'Trae IDE Agent Skills Case Study', 'Automating Onboarding & Management Workflows')

    # 2. Project Overview
    add_bullet_slide(prs, 'Project Overview', [
        {'text': 'Objective', 'bold': True},
        {'text': 'Streamline new employee onboarding, organizational management, and team enablement using AI Agents.', 'level': 1},
        {'text': 'Key Technologies', 'bold': True},
        {'text': 'Trae IDE, Claude Agent Skills, Python (Pandas), HTML/JS/Tailwind.', 'level': 1}
    ])

    # 3. Skill 1: Intelligent Excel Processing
    add_bullet_slide(prs, 'Skill 1: Intelligent Excel Processing', [
        {'text': 'Skill Used: xlsx / RunCommand', 'bold': True},
        {'text': 'Challenge: Managing complex organizational charts and training data.', 'level': 0},
        {'text': 'Actions Taken:', 'level': 0},
        {'text': 'Analyzed "中欧总表.xlsx" to extract trainee progress (Zou Fengjing: 4%).', 'level': 1},
        {'text': 'Batch replaced sensitive personnel names with "Pending" (待定) using Regex.', 'level': 1},
        {'text': 'Converted tabular data into Mermaid structure diagrams.', 'level': 1},
        {'text': 'Outcome: Automated data hygiene and visualization preparation.', 'level': 1}
    ])

    # 4. Skill 2: Workflow Automation
    add_bullet_slide(prs, 'Skill 2: Workflow Automation', [
        {'text': 'Skill Used: outlook-sender', 'bold': True},
        {'text': 'Challenge: Manual email notifications for reports.', 'level': 0},
        {'text': 'Actions Taken:', 'level': 0},
        {'text': 'Created a custom skill to interface with local Outlook via PowerShell.', 'level': 1},
        {'text': 'Designed scripts for sending emails with attachments and saving drafts.', 'level': 1},
        {'text': 'Outcome: Standardized communication capability (demonstrated proof-of-concept).', 'level': 1}
    ])

    # 5. Skill 3: Frontend Data Visualization
    add_bullet_slide(prs, 'Skill 3: Frontend Data Visualization', [
        {'text': 'Skill Used: frontend-design', 'bold': True},
        {'text': 'Challenge: Text-based status reports are hard to digest.', 'level': 0},
        {'text': 'Actions Taken:', 'level': 0},
        {'text': 'Transformed Markdown status reports into an interactive HTML dashboard.', 'level': 1},
        {'text': 'Implemented real-time progress charts using Chart.js.', 'level': 1},
        {'text': 'Added interactive features: Role Detection & Weight Adjustment.', 'level': 1},
        {'text': 'Outcome: Professional, responsive onboarding status board.', 'level': 1}
    ])

    # 6. Skill 4: Knowledge Enablement
    add_bullet_slide(prs, 'Skill 4: Knowledge Enablement', [
        {'text': 'Skill Used: enjoy-bot-community-guide', 'bold': True},
        {'text': 'Challenge: Team lacks best practices for using AI community resources.', 'level': 0},
        {'text': 'Actions Taken:', 'level': 0},
        {'text': 'Encapsulated community guidelines into a retrievable Skill.', 'level': 1},
        {'text': 'Defined "A-B-C" usage strategy (Scenarios, Structure, Action).', 'level': 1},
        {'text': 'Standardized "High-Quality Question Template".', 'level': 1},
        {'text': 'Outcome: Scalable knowledge transfer mechanism.', 'level': 1}
    ])

    # 7. Summary
    add_bullet_slide(prs, 'Conclusion', [
        {'text': 'Impact of Agent Skills:', 'bold': True},
        {'text': 'Structured: Converting ad-hoc tasks into reusable modules.', 'level': 1},
        {'text': 'Automated: Reducing manual effort in data processing and reporting.', 'level': 1},
        {'text': 'Interactive: Moving from static text to dynamic visual tools.', 'level': 1},
        {'text': 'This project demonstrates how Trae IDE Agents can act as full-stack engineers to handle end-to-end workflows.', 'level': 0}
    ])

    prs.save('Trae_Agent_Skills_Case_Study.pptx')
    print('Presentation saved successfully.')

create_presentation()
